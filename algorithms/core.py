"""Shared core business, parsing, and plotting logic for LaczTabeleNietoperzy.

Single canonical source of truth for both Tkinter desktop and headless CLI workflows.
100% offline with zero network dependencies or HTTP imports.
"""

from __future__ import annotations

import math
import os
import re
import subprocess
import sys
from datetime import date, datetime, time, timedelta
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from matplotlib.ticker import MaxNLocator
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

try:
    from algorithms.input_reader import read_input_table
except ModuleNotFoundError:
    from input_reader import read_input_table

# ================== Grundinställningar / etiketter ==================
DEFAULT_OUT_BASENAME = "sammanstallning_fladdermus"

# ================== Färger – enhetligt HEX-format ==================
HEX_HEADER_BG = "#595959"
HEX_HEADER_FG = "#FFFFFF"
HEX_TABLE_FORBI = "#E7E6E6"

# Standardfärger när INGEN SOF finns i datasetet
HEX_NVI_SOC_DEFAULT_NO_SOF = "#FF0000"
HEX_ART_SOC_DEFAULT_NO_SOF = "#EB09D8"

# Standardfärger när SOF FINNS i datasetet (#8)
HEX_NVI_SOC_DEFAULT_WITH_SOF = "#9A0B08"  # SOC = socialt - läte
HEX_NVI_SOF_DEFAULT          = "#FF0000"  # SOF = socialt - flyg (tar ursprungliga SOC-färgen)

HEX_ART_SOC_DEFAULT_WITH_SOF = "#8D0B82"  # SOC = socialt - läte
HEX_ART_SOF_DEFAULT          = "#EB09D8"  # SOF = socialt - flyg (tar ursprungliga SOC-färgen)

HEX_NVI_FODO  = "#FFC000"  # Födosökande
HEX_NVI_FORBI = "#A9A9A9"  # Förbiflygande

HEX_ART_FODO  = "#D98FD3"  # Födosökande
HEX_ART_FORBI = "#ABAAA9"  # Förbiflygande

HEX_NVI_SOC = HEX_NVI_SOC_DEFAULT_NO_SOF
HEX_ART_SOC = HEX_ART_SOC_DEFAULT_NO_SOF

BORDER_MEDIUM = Side(style="medium", color="FF000000")


def hex_to_argb(hex_rgb: str) -> str:
    h = hex_rgb.strip().lstrip("#")
    if len(h) != 6:
        raise ValueError(f"Ogiltig HEX: {hex_rgb}")
    return "FF" + h.upper()


def fill_from_hex(hex_rgb: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_to_argb(hex_rgb))


FILL_HDR = fill_from_hex(HEX_HEADER_BG)


def report_progress(pct: int, msg: str) -> None:
    progress_path_str = os.environ.get("PROGRESS_PATH")
    if progress_path_str:
        try:
            import json

            p = os.path.abspath(progress_path_str)
            os.makedirs(os.path.dirname(p), exist_ok=True)
            tmp_p = f"{p}.tmp"
            with open(tmp_p, "w", encoding="utf-8") as f:
                json.dump({"progress": pct, "progress_message": msg}, f)
            os.replace(tmp_p, p)
        except Exception:
            pass


# =============== Ordbok Latin → Svenska ===============
LATIN_TO_SV = {
    "Barbastella barbastellus": "Barbastell",
    "Eptesicus nilssonii": "Nordfladdermus",
    "Eptesicus serotinus": "Sydfladdermus",
    "Cnephaeus nilssonii": "Nordfladdermus",
    "Cnephaeus serotinus": "Sydfladdermus",
    "Myotis alcathoe": "Nymffladdermus",
    "Myotis bechsteinii": "Bechsteins fladdermus",
    "Myotis brandtii": "Tajgafladdermus",
    "Myotis dasycneme": "Dammfladdermus",
    "Myotis daubentonii": "Vattenfladdermus",
    "Myotis myotis": "Större musöra",
    "Myotis mystacinus": "Mustaschfladdermus",
    "Myotis nattereri": "Fransfladdermus",
    "Myotis mystacinus/brandtii": "Mustasch/Tajgafladdermus",
    "Nyctalus leisleri": "Mindre brunfladdermus",
    "Nyctalus noctula": "Större brunfladdermus",
    "Pipistrellus kuhlii": "Parkpipistrell",
    "Pipistrellus nathusii": "Trollpipistrell",
    "Pipistrellus pipistrellus": "Sydpipistrell",
    "Pipistrellus pygmaeus": "Dvärgpipistrell",
    "Plecotus auritus": "Brunlångöra",
    "Plecotus austriacus": "Grålångöra",
    "Vespertilio murinus": "Gråskimlig fladdermus",
    "Nyctaloid": None,
    "Chiroptera": None,
}
SPECIAL_TAIL = {"nyctaloid", "chiroptera"}

CODE_TO_LATIN = {
    "NYCNOC": "Nyctalus noctula",
    "PLEUAR": "Plecotus auritus",
    "VESMUR": "Vespertilio murinus",
    "PIPNAT": "Pipistrellus nathusii",
    "PIPPYG": "Pipistrellus pygmaeus",
    "PIPKUH": "Pipistrellus kuhlii",
    "PIPPIP": "Pipistrellus pipistrellus",
    "EPTNIL": "Eptesicus nilssonii",
    "EPTSER": "Eptesicus serotinus",
    "MYODAB": "Myotis daubentonii",
    "MYOMYS": "Myotis mystacinus",
    "MYONAT": "Myotis nattereri",
    "MYODAS": "Myotis dasycneme",
    "MYOBEC": "Myotis bechsteinii",
    "MYOALC": "Myotis alcathoe",
    "MYOBRA": "Myotis brandtii",
    "BARBAR": "Barbastella barbastellus",
    "NYCLEI": "Nyctalus leisleri",
}


def get_swedish_species_name(sp: str) -> str:
    sp_str = str(sp).strip()
    if not sp_str:
        return sp_str
    if sp_str in LATIN_TO_SV and LATIN_TO_SV[sp_str]:
        return LATIN_TO_SV[sp_str]
    latin = CODE_TO_LATIN.get(sp_str.upper(), sp_str)
    sv = LATIN_TO_SV.get(latin)
    if sv:
        return sv
    return sp_str


def safe_sheet_name(path: str, used: set) -> str:
    base = os.path.splitext(os.path.basename(path))[0]
    base = re.sub(r'[:\\\/\?\*\[\]]', '_', base).strip()[:31] or "Ark"
    cand = base
    i = 2
    while cand in used or not cand:
        suf = f"_{i}"
        cand = (base[: (31 - len(suf))] + suf) if len(base) + len(suf) > 31 else (base + suf)
        i += 1
    used.add(cand)
    return cand


def display_label_multiline(latin: str) -> str:
    latin = str(latin).strip()
    sv = LATIN_TO_SV.get(latin, None)
    if sv:
        return f"{sv.capitalize()},\n{latin}"
    return latin


def format_title(species_latin: str, total_count: int, night_label: str | None = None) -> str:
    sv = LATIN_TO_SV.get(species_latin)
    display_latin = species_latin
    if species_latin == "Eptesicus nilssonii":
        display_latin = "Cnephaeus nilssonii"
    elif species_latin == "Eptesicus serotinus":
        display_latin = "Cnephaeus serotinus"

    base = f"{sv} ({display_latin})" if sv else display_latin
    if night_label:
        return f"{base} – natt {night_label}, antal observerade beteenden: {int(total_count)}"
    return f"{base}, antal observerade beteenden: {int(total_count)}"


def extract_species_and_type(manual_id_value):
    """Parse MANUAL ID column into list of (species, type_code) tuples."""
    if pd.isna(manual_id_value):
        return []
    out = []
    for raw in str(manual_id_value).split(","):
        entry = raw.strip()
        if not entry:
            continue
        upper = entry.upper()
        if "SOF" in upper or "SOCIALT - FLYG" in upper:
            typ = "SOF"
        elif "SOC" in upper or "SOCIALT" in upper:
            typ = "SOC"
        elif "FOD" in upper or "FÖDOSÖKANDE" in upper or "FODOSOKANDE" in upper:
            typ = "FOD"
        else:
            typ = "FORBI"

        species = re.sub(r'\b(FOD|SOC|SOF|FORBI)\b', "", entry, flags=re.IGNORECASE)
        species = re.sub(
            r'(?i)socialt\s*-\s*flyg|socialt\s*-\s*läte|socialt|födosökande|fodosokande|förbiflygande',
            "",
            species,
        ).strip(" ,;")
        species = CODE_TO_LATIN.get(species.upper(), species)
        if species == "Eptesicus nilssonii":
            species = "Cnephaeus nilssonii"
        elif species == "Eptesicus serotinus":
            species = "Cnephaeus serotinus"
        out.append((species, typ))
    return out


def resolve_dataset_class_config(input_files_or_dfs, custom_colors=None):
    """Determine if entire dataset contains SOF and build class names, type order, and color palettes."""
    has_sof = False
    for item in input_files_or_dfs:
        if isinstance(item, (str, os.PathLike)):
            try:
                df = read_input_table(item)
            except Exception:
                continue
        else:
            df = item

        if "MANUAL ID" in df.columns:
            for val in df["MANUAL ID"].dropna():
                for _, typ in extract_species_and_type(val):
                    if typ == "SOF":
                        has_sof = True
                        break
                if has_sof:
                    break
        if has_sof:
            break

    custom_colors = custom_colors or {}
    nvi_custom = custom_colors.get("NVI") or {}
    art_custom = custom_colors.get("ART") or {}

    if has_sof:
        type_order = ["Socialt - läte", "Socialt - flyg", "Födosökande", "Förbiflygande"]
        code_map = {
            "SOC": "Socialt - läte",
            "SOF": "Socialt - flyg",
            "FOD": "Födosökande",
            "FORBI": "Förbiflygande",
        }
        nvi_soc_color = (
            nvi_custom.get("Socialt - läte")
            or nvi_custom.get("SOC")
            or nvi_custom.get("Socialt")
            or HEX_NVI_SOC_DEFAULT_WITH_SOF
        )
        nvi_sof_color = nvi_custom.get("Socialt - flyg") or nvi_custom.get("SOF") or HEX_NVI_SOF_DEFAULT
        nvi_fodo_color = nvi_custom.get("Födosökande") or nvi_custom.get("FOD") or HEX_NVI_FODO
        nvi_forbi_color = nvi_custom.get("Förbiflygande") or nvi_custom.get("FORBI") or HEX_NVI_FORBI

        art_soc_color = (
            art_custom.get("Socialt - läte")
            or art_custom.get("SOC")
            or art_custom.get("Socialt")
            or HEX_ART_SOC_DEFAULT_WITH_SOF
        )
        art_sof_color = art_custom.get("Socialt - flyg") or art_custom.get("SOF") or HEX_ART_SOF_DEFAULT
        art_fodo_color = art_custom.get("Födosökande") or art_custom.get("FOD") or HEX_ART_FODO
        art_forbi_color = art_custom.get("Förbiflygande") or art_custom.get("FORBI") or HEX_ART_FORBI

        nvi_colors = [nvi_soc_color, nvi_sof_color, nvi_fodo_color, nvi_forbi_color]
        art_colors = [art_soc_color, art_sof_color, art_fodo_color, art_forbi_color]
    else:
        type_order = ["Socialt", "Födosökande", "Förbiflygande"]
        code_map = {
            "SOC": "Socialt",
            "SOF": "Socialt - flyg",
            "FOD": "Födosökande",
            "FORBI": "Förbiflygande",
        }
        nvi_soc_color = (
            nvi_custom.get("Socialt")
            or nvi_custom.get("Socialt - läte")
            or nvi_custom.get("SOC")
            or HEX_NVI_SOC_DEFAULT_NO_SOF
        )
        nvi_fodo_color = nvi_custom.get("Födosökande") or nvi_custom.get("FOD") or HEX_NVI_FODO
        nvi_forbi_color = nvi_custom.get("Förbiflygande") or nvi_custom.get("FORBI") or HEX_NVI_FORBI

        art_soc_color = (
            art_custom.get("Socialt")
            or art_custom.get("Socialt - läte")
            or art_custom.get("SOC")
            or HEX_ART_SOC_DEFAULT_NO_SOF
        )
        art_fodo_color = art_custom.get("Födosökande") or art_custom.get("FOD") or HEX_ART_FODO
        art_forbi_color = art_custom.get("Förbiflygande") or art_custom.get("FORBI") or HEX_ART_FORBI

        nvi_colors = [nvi_soc_color, nvi_fodo_color, nvi_forbi_color]
        art_colors = [art_soc_color, art_fodo_color, art_forbi_color]

    nvi_color_dict = dict(zip(type_order, nvi_colors))
    art_color_dict = dict(zip(type_order, art_colors))

    return {
        "has_sof": has_sof,
        "type_order": type_order,
        "code_map": code_map,
        "colors_art": art_colors,
        "colors_nvi": nvi_colors,
        "art_color_dict": art_color_dict,
        "nvi_color_dict": nvi_color_dict,
    }


def compute_ymax_with_headroom(max_val: float | int, headroom_factor: float = 1.05) -> int:
    """Derive Y max from data with headroom.
    
    Fixed mode: headroom_factor = 1.05 (legacy 5% headroom).
    Zoomed mode: headroom_factor = 1.10 (zoomed 10% headroom).
    Deterministic min of 1 when max_val <= 0.
    """
    if max_val <= 0:
        return 1
    return max(1, math.ceil(max_val * headroom_factor))


def open_file(path):
    try:
        if sys.platform.startswith("win"):
            os.startfile(path)
        elif sys.platform == "darwin":
            subprocess.run(["open", path])
        else:
            subprocess.run(["xdg-open", path])
    except Exception as e:
        print(f"Kan inte öppna filen automatiskt: {e}")


def species_sort_key(latin: str):
    return (str(latin).strip().lower() in SPECIAL_TAIL, str(latin).casefold())


def safe_filename(s):
    return re.sub(r'[\\/:\*\?"<>\|]', '_', str(s))


# --- Tidshjälp ---
def validate_hhmm(val: str | None) -> tuple[int, int] | None:
    if val is None or not str(val).strip():
        return None
    s = str(val).strip()
    match = re.fullmatch(r"([0-1][0-9]|2[0-3]):([0-5][0-9])", s)
    if not match:
        raise ValueError(f"Invalid time format '{val}'. Expected HH:MM in 24-hour format.")
    return int(match.group(1)), int(match.group(2))


def is_time_in_range(time_val, start_str: str, stop_str: str) -> bool:
    hm = _hm_from_any(time_val)
    if hm is None:
        return False
    sh_sm = validate_hhmm(start_str)
    eh_em = validate_hhmm(stop_str)
    if not sh_sm or not eh_em:
        return True
    t_min = time(sh_sm[0], sh_sm[1])
    t_max = time(eh_em[0], eh_em[1])
    t_reg = time(hm[0], hm[1])
    if sh_sm <= eh_em:
        return t_min <= t_reg <= t_max
    else:
        return t_reg >= t_min or t_reg <= t_max


def build_manual_interval_sequence(custom_time_range: tuple[str, str]) -> list[str]:
    start_str, stop_str = custom_time_range
    sh_sm = validate_hhmm(start_str)
    eh_em = validate_hhmm(stop_str)
    if not sh_sm or not eh_em:
        return []
    start_dt = datetime(2000, 1, 1, sh_sm[0], sh_sm[1])
    if eh_em < sh_sm:
        stop_dt = datetime(2000, 1, 2, eh_em[0], eh_em[1])
    else:
        stop_dt = datetime(2000, 1, 1, eh_em[0], eh_em[1])
    min_dt = round_down_15(start_dt)
    max_dt = round_up_15(stop_dt)

    intervals = []
    t = min_dt
    while t <= max_dt:
        intervals.append(t.strftime("%H:%M"))
        t += timedelta(minutes=15)
    return list(dict.fromkeys(intervals))


def get_unique_input_stems(input_files: list[str]) -> dict[str, str]:
    used_stems = set()
    result = {}
    for path in input_files:
        base_stem = safe_filename(os.path.splitext(os.path.basename(path))[0])
        cand = base_stem
        counter = 2
        while cand in used_stems:
            cand = f"{base_stem}_{counter}"
            counter += 1
        used_stems.add(cand)
        result[path] = cand
    return result


def _hm_from_any(val):
    try:
        if val is None or (isinstance(val, float) and pd.isna(val)) or (isinstance(val, str) and val.strip() == ""):
            return None
        if hasattr(val, "hour") and hasattr(val, "minute"):
            return int(val.hour), int(val.minute)
        if isinstance(val, (int, float)) and not pd.isna(val):
            frac = float(val) % 1.0
            secs = int(round(frac * 24 * 60 * 60))
            h = (secs // 3600) % 24
            m = (secs % 3600) // 60
            return int(h), int(m)
        s = str(val).strip()
        t = pd.to_datetime(s, format="%H:%M:%S", errors="coerce")
        if pd.isna(t):
            t = pd.to_datetime(s, format="%H:%M", errors="coerce")
        if pd.isna(t):
            return None
        return int(t.hour), int(t.minute)
    except Exception:
        return None


def str_to_dt(time_val):
    hm = _hm_from_any(time_val)
    if hm is None:
        return None
    h, m = hm
    fake_date = "2000-01-02" if h < 12 else "2000-01-01"
    return datetime.strptime(f"{fake_date} {h:02d}:{m:02d}", "%Y-%m-%d %H:%M")


def round_down_15(dt):
    return dt.replace(minute=(dt.minute // 15) * 15, second=0, microsecond=0)


def round_up_15(dt):
    if dt.minute % 15 != 0 or dt.second > 0 or dt.microsecond > 0:
        dt = dt + timedelta(minutes=15 - (dt.minute % 15), seconds=-dt.second, microseconds=-dt.microsecond)
    return dt.replace(second=0, microsecond=0)


def interval_to_sortkey(interval):
    try:
        t = pd.to_datetime(str(interval), format="%H:%M", errors="coerce")
        if pd.isna(t):
            return None
        h, m = int(t.hour), int(t.minute)
        fake_date = "2000-01-02" if h < 12 else "2000-01-01"
        return datetime.strptime(f"{fake_date} {h:02d}:{m:02d}", "%Y-%m-%d %H:%M")
    except Exception:
        return None


def detect_column(df, candidates):
    lowmap = {str(c).strip().lower(): c for c in df.columns}
    for k in candidates:
        if k in lowmap:
            return lowmap[k]
    return None


def count_nights(df):
    date_col = detect_column(df, ["date", "datum"])
    if not date_col:
        return None
    time_col = detect_column(df, ["time", "tid"])

    dates = pd.to_datetime(df[date_col], errors="coerce")
    if time_col:
        hm = df[time_col].apply(_hm_from_any)
        hours = hm.apply(lambda x: x[0] if isinstance(x, tuple) else None)
        shift = hours.apply(lambda h: (h is not None) and (h < 12))
        night_key = (dates.dt.normalize() - pd.to_timedelta(shift.fillna(False).astype(int), unit="D")).dt.date
    else:
        night_key = dates.dt.normalize().dt.date

    nights = pd.Series(night_key).dropna().nunique()
    return int(nights) if nights else None


def row_night_key(d_val, t_val):
    d = pd.to_datetime(d_val, errors="coerce")
    if pd.isna(d):
        return None
    hm = _hm_from_any(t_val)
    if hm is None:
        return d.date()
    h, _ = hm
    return (d - pd.Timedelta(days=1)).date() if h < 12 else d.date()


def night_label_str(night_start: date) -> str:
    nxt = night_start + timedelta(days=1)
    return f"{night_start.day}/{nxt.day}.{night_start.month:02d}"


def _validate_hex(s):
    s = (s or "").strip()
    if not s:
        return None
    if s.startswith("#"):
        s = s[1:]
    if len(s) != 6 or any(c not in "0123456789abcdefABCDEF" for c in s):
        return None
    return "#" + s.upper()


# ================== Y-Max Computation ==================
def _compute_ymax_for_subset(df_subset, custom_time_range, code_map=None):
    if df_subset.empty:
        return 0
    df = df_subset.copy()
    if "MANUAL ID" not in df.columns:
        return 0

    time_col = detect_column(df, ["time", "tid"])
    if time_col is None:
        return 0

    if custom_time_range:
        start_str, stop_str = custom_time_range
        df = df[df[time_col].apply(lambda t: is_time_in_range(t, start_str, stop_str))]
        if df.empty:
            return 0

    df["species_type_list"] = df["MANUAL ID"].map(extract_species_and_type)

    def time_to_interval(val):
        hm = _hm_from_any(val)
        if hm is None:
            return ""
        h, m = hm
        minutes = int((m // 15) * 15)
        return f"{h:02d}:{minutes:02d}"

    df["interval"] = df[time_col].map(time_to_interval)

    df_long = df.explode("species_type_list")
    df_long = df_long[df_long["species_type_list"].notna()]
    if df_long.empty:
        return 0
    df_long[["species", "obs_code"]] = pd.DataFrame(df_long["species_type_list"].tolist(), index=df_long.index)
    if code_map:
        df_long["obs_type"] = df_long["obs_code"].map(code_map)
    else:
        df_long["obs_type"] = df_long["obs_code"]

    df_long = df_long[df_long["species"].astype(str).str.strip().str.lower() != "noise"]
    df_long = df_long[df_long["species"].astype(str).str.strip() != ""]
    if df_long.empty:
        return 0

    if custom_time_range:
        all_intervals = build_manual_interval_sequence(custom_time_range)
    else:
        dt_series = df[time_col].apply(str_to_dt).dropna()
        if len(dt_series) == 0:
            ints = [s for s in df["interval"].astype(str).tolist() if s and s.lower() != "nan"]
            dt_from_int = [interval_to_sortkey(s) for s in ints]
            dt_from_int = [d for d in dt_from_int if d is not None]
            if not dt_from_int:
                return 0
            min_dt = round_down_15(min(dt_from_int))
            max_dt = round_up_15(max(dt_from_int))
        else:
            min_dt = round_down_15(min(dt_series))
            max_dt = round_up_15(max(dt_series))

        all_intervals = []
        t = min_dt
        while t <= max_dt:
            all_intervals.append(t.strftime("%H:%M"))
            t += timedelta(minutes=15)
        all_intervals = list(dict.fromkeys(all_intervals))

    agg = df_long.groupby(["interval", "species", "obs_type"]).size().reset_index(name="antal")

    y_max = 0
    for sp in df_long["species"].unique():
        plot_data = (
            agg[agg["species"] == sp]
            .pivot(index="interval", columns="obs_type", values="antal")
            .fillna(0)
            .reindex(all_intervals, fill_value=0)
        )
        if not plot_data.empty:
            y_max = max(y_max, int(plot_data.sum(axis=1).max()))
    return y_max


def compute_global_ymax_across_files(input_files_list, custom_time_range, code_map=None):
    """Legacy global Y-max computation using 5% headroom factor (fixed mode)."""
    y_global = 0
    for path in input_files_list:
        df = read_input_table(path)
        y_global = max(y_global, _compute_ymax_for_subset(df, custom_time_range, code_map=code_map))
    return compute_ymax_with_headroom(y_global, headroom_factor=1.05)


def compute_global_ymax_across_files_and_nights(input_files_list, custom_time_range, code_map=None):
    """Legacy global Y-max computation across files and nights using 5% headroom factor (fixed mode)."""
    y_global = 0
    for path in input_files_list:
        df_full = read_input_table(path)
        date_col = detect_column(df_full, ["date", "datum"])
        time_col = detect_column(df_full, ["time", "tid"])
        if not date_col:
            y_global = max(y_global, _compute_ymax_for_subset(df_full, custom_time_range, code_map=code_map))
            continue
        df_full["__night"] = df_full.apply(
            lambda r: row_night_key(r[date_col], r[time_col] if time_col in r else None), axis=1
        )
        for night_key, sub in df_full.groupby("__night"):
            if night_key is None:
                continue
            y_global = max(y_global, _compute_ymax_for_subset(sub, custom_time_range, code_map=code_map))
    return compute_ymax_with_headroom(y_global, headroom_factor=1.05)


# ================== #9 All-Species Grouped Data Construction ==================
def build_all_species_grouped_data(
    df_long: pd.DataFrame,
    all_intervals: list[str],
    type_order: list[str],
) -> tuple[list[dict], int]:
    """Construct ordered structured data containing ONLY species present inside each time bin (#9).

    Grouping order:
      1. Time bin (chronological sequence of all_intervals)
      2. Species present within that specific time bin (sorted deterministically)
      3. Class breakdown per species (type_order)

    Returns:
      (grouped_items, peak_bar_height)
      where grouped_items is a list of dicts:
        {
            "interval": str,
            "species": str,
            "class_counts": dict[str, int],
            "total_count": int,
        }
    """
    if df_long.empty or not all_intervals:
        return [], 0

    df_filtered = df_long[df_long["interval"].isin(all_intervals)]
    if df_filtered.empty:
        return [], 0

    grouped_items = []
    peak_bar_height = 0

    for intv in all_intervals:
        df_intv = df_filtered[df_filtered["interval"] == intv]
        if df_intv.empty:
            continue

        present_species = sorted(df_intv["species"].unique(), key=species_sort_key)

        for sp in present_species:
            df_sp = df_intv[df_intv["species"] == sp]
            class_counts = {}
            for ot in type_order:
                cnt = int((df_sp["obs_type"] == ot).sum())
                if cnt > 0:
                    class_counts[ot] = cnt

            bar_h = sum(class_counts.values())
            if bar_h > peak_bar_height:
                peak_bar_height = bar_h

            grouped_items.append({
                "interval": str(intv),
                "species": str(sp),
                "class_counts": class_counts,
                "total_count": bar_h,
            })

    return grouped_items, peak_bar_height


def compute_all_species_geometry(max_k: int, num_intervals: int, scale: float = 1.0) -> dict[str, float | int]:
    """Compute unified all-species chart geometry derived from one shared base bar width (#9).

    Reference Geometry Principle:
      base_bar_width = 0.16 * scale
      species_spacing = base_bar_width
      slot_width = (max_k + 0.8) * base_bar_width

      fig_scale = 0.11 / base_bar_width
      width_per_interval = max(0.20, slot_width * fig_scale)
      fig_w = max(8.0, min(35.0, num_intervals * width_per_interval + 2.5))
    """
    k_eff = max(1, int(max_k))
    base_bar_width = 0.16 * scale
    species_spacing = base_bar_width
    slot_width = (k_eff + 0.8) * base_bar_width

    fig_scale = 0.11 / base_bar_width
    width_per_interval = max(0.20, slot_width * fig_scale)
    fig_w = max(8.0, min(35.0, float(num_intervals * width_per_interval + 2.5)))

    return {
        "max_k": k_eff,
        "base_bar_width": base_bar_width,
        "species_spacing": species_spacing,
        "slot_width": slot_width,
        "fig_w": fig_w,
    }


def _plot_all_species_grouped_stacked(
    df_long: pd.DataFrame,
    all_intervals: list[str],
    species_list: list[str],
    type_order: list[str],
    color_dict: dict[str, str],
    out_path: str,
    title_text: str,
    ymax_mode: str,
    y_lim_global: int,
):
    """Plot grouped stacked bar chart for all species across time intervals (#9).
    
    Layout specifications:
    - Dual horizontal axes:
        * Top axis (`ax_top`): 15-minute time labels centered over equal-width slots.
        * Bottom axis (`ax`): Vertical species labels directly under rendered species bars.
    - Subtle vertical dotted separators between adjacent 15-minute time slots.
    - Stacked bar count annotation centered above each non-zero species stack.
    - Equal-width 15-minute slot geometry derived from single shared reference geometry.
    """
    if not all_intervals:
        return

    grouped_items, peak_bar_h = build_all_species_grouped_data(df_long, all_intervals, type_order)

    if ymax_mode == "zoomed":
        chart_ymax = compute_ymax_with_headroom(peak_bar_h, headroom_factor=1.10)
    else:
        chart_ymax = y_lim_global

    if chart_ymax <= 0:
        chart_ymax = 1

    num_intervals = len(all_intervals)
    items_by_interval: dict[str, list[dict]] = {}
    for item in grouped_items:
        items_by_interval.setdefault(item["interval"], []).append(item)

    max_k = max([len(items) for items in items_by_interval.values()], default=1)
    geom = compute_all_species_geometry(max_k, num_intervals)

    slot_width = float(geom["slot_width"])
    base_bar_width = float(geom["base_bar_width"])
    species_spacing = float(geom["species_spacing"])
    fig_w = float(geom["fig_w"])

    fig, ax = plt.subplots(figsize=(fig_w, 7))
    ax_top = ax.twiny()

    legend_handles = {}

    top_tick_positions = []
    top_tick_labels = []

    species_tick_positions = []
    species_tick_labels = []

    for idx, intv in enumerate(all_intervals):
        slot_center = idx * slot_width
        top_tick_positions.append(slot_center)
        top_tick_labels.append(str(intv))

        # Subtle vertical dotted separator between adjacent time slots
        if idx < num_intervals - 1:
            x_sep = slot_center + slot_width / 2.0
            ax.axvline(x=x_sep, color="#cccccc", linestyle=":", linewidth=0.8, zorder=0)

        intv_items = items_by_interval.get(intv, [])
        k = len(intv_items)
        if k == 0:
            continue

        for j, item in enumerate(intv_items):
            x_pos = slot_center + (j - (k - 1) / 2.0) * species_spacing
            sp = item["species"]
            class_counts = item["class_counts"]

            bottom = 0
            for ot in type_order:
                val = class_counts.get(ot, 0)
                color = color_dict.get(ot, "#000000")
                if val > 0:
                    bar_container = ax.bar(
                        x_pos,
                        val,
                        bottom=bottom,
                        width=base_bar_width,
                        color=color,
                        edgecolor="none",
                    )
                    if ot not in legend_handles:
                        legend_handles[ot] = bar_container[0]
                    bottom += val

            total_count = item["total_count"]
            if total_count > 0:
                ax.text(
                    x_pos,
                    bottom + chart_ymax * 0.01,
                    str(total_count),
                    ha="center",
                    va="bottom",
                    fontsize=7,
                    fontweight="bold",
                )

            sp_label = get_swedish_species_name(sp)
            species_tick_positions.append(x_pos)
            species_tick_labels.append(sp_label)

    # Set up axes geometry and limits
    x_min = -slot_width / 2.0
    x_max = num_intervals * slot_width - slot_width / 2.0
    ax.set_xlim(x_min, x_max)
    ax_top.set_xlim(x_min, x_max)

    # Top axis (15-minute time intervals)
    ax_top.set_xticks(top_tick_positions)
    ax_top.set_xticklabels(top_tick_labels, rotation=90)
    ax_top.set_xlabel("Tid (15-minutersintervall)", labelpad=10)
    ax_top.set_title(title_text, pad=50)

    # Bottom axis (species)
    ax.set_xticks(species_tick_positions)
    ax.set_xticklabels(species_tick_labels, rotation=90)
    ax.set_ylabel("Antal ljudfiler")

    if legend_handles:
        handles = [legend_handles[ot] for ot in type_order if ot in legend_handles]
        labels = [ot for ot in type_order if ot in legend_handles]
        ax.legend(handles, labels, title="Beteendetyper", bbox_to_anchor=(1.02, 1), loc="upper left")

    ax.yaxis.set_major_locator(MaxNLocator(integer=True))
    ax.set_ylim(0, chart_ymax)
    ax.grid(True, axis="y")
    plt.tight_layout()
    plt.savefig(out_path)
    plt.close()


