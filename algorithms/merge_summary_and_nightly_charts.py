"""Merge bat Excel tables and create combined and/or per-night charts.

GUI udostępnia dwie niezależne opcje: wykresy zbiorcze oraz noc-po-nocy.
Wykresy nocne są tworzone dodatkowo, z etykietą DD/NN.MM w tytułach.
Eksport Excel, globalna oś Y, kolory i zakres osi X pozostają wspólne.
"""

import os
import re
import sys
import math
import subprocess
import pandas as pd
import matplotlib.pyplot as plt
from datetime import datetime, timedelta, date, time
from matplotlib.ticker import MaxNLocator

import tkinter as tk
from tkinter import filedialog, messagebox, ttk, colorchooser

from openpyxl import load_workbook
from openpyxl.styles import Alignment, PatternFill, Font, Border, Side
from openpyxl.utils import get_column_letter

try:
    from algorithms.input_reader import read_input_table
except ModuleNotFoundError:  # Direct execution from the algorithms directory.
    from input_reader import read_input_table

try:
    import algorithms.core as core
    from algorithms.core import (
        BORDER_MEDIUM,
        DEFAULT_OUT_BASENAME,
        FILL_HDR,
        HEX_HEADER_FG,
        HEX_NVI_FODO,
        HEX_ART_FODO,
        HEX_TABLE_FORBI,
        LATIN_TO_SV,
        _compute_ymax_for_subset,
        _hm_from_any,
        _validate_hex,
        build_all_species_grouped_data,
        build_manual_interval_sequence,
        compute_global_ymax_across_files,
        compute_global_ymax_across_files_and_nights,
        compute_ymax_with_headroom,
        count_nights,
        detect_column,
        display_label_multiline,
        extract_species_and_type,
        fill_from_hex,
        format_title,
        get_unique_input_stems,
        hex_to_argb,
        is_time_in_range,
        interval_to_sortkey,
        night_label_str,
        open_file,
        report_progress,
        resolve_dataset_class_config,
        round_down_15,
        round_up_15,
        row_night_key,
        safe_filename,
        safe_sheet_name,
        species_sort_key,
        str_to_dt,
        validate_hhmm,
    )
except ModuleNotFoundError:
    import core
    from core import (
        BORDER_MEDIUM,
        DEFAULT_OUT_BASENAME,
        FILL_HDR,
        HEX_HEADER_FG,
        HEX_NVI_FODO,
        HEX_ART_FODO,
        HEX_TABLE_FORBI,
        LATIN_TO_SV,
        _compute_ymax_for_subset,
        _hm_from_any,
        _validate_hex,
        build_all_species_grouped_data,
        build_manual_interval_sequence,
        compute_global_ymax_across_files,
        compute_global_ymax_across_files_and_nights,
        compute_ymax_with_headroom,
        count_nights,
        detect_column,
        display_label_multiline,
        extract_species_and_type,
        fill_from_hex,
        format_title,
        get_unique_input_stems,
        hex_to_argb,
        is_time_in_range,
        interval_to_sortkey,
        night_label_str,
        open_file,
        report_progress,
        resolve_dataset_class_config,
        round_down_15,
        round_up_15,
        row_night_key,
        safe_filename,
        safe_sheet_name,
        species_sort_key,
        str_to_dt,
        validate_hhmm,
    )


def gui_collect_settings(default_basename=DEFAULT_OUT_BASENAME):
    root = tk.Tk()
    root.title("Fladdermus – sammanställning & diagram (natt/löpande)")
    root.geometry("920x840"); root.minsize(860, 700)
    main = ttk.Frame(root, padding=12); main.pack(fill="both", expand=True)

    # Indata
    lf_in = ttk.LabelFrame(main, text="Indatafiler"); lf_in.pack(fill="both", expand=False, pady=(0,10))
    files_list = tk.Listbox(lf_in, height=6, selectmode="extended")
    files_list.grid(row=0, column=0, rowspan=3, sticky="nsew", padx=(8,8), pady=8)
    lf_in.columnconfigure(0, weight=1); lf_in.rowconfigure(0, weight=1)
    ttk.Button(lf_in, text="Lägg till filer…",
               command=lambda: [files_list.insert(tk.END, p) for p in filedialog.askopenfilenames(
                   title="Välj XLSX- eller CSV-filer", filetypes=[
                       ("XLSX- och CSV-filer", "*.xlsx *.csv"),
                       ("Excel", "*.xlsx"),
                       ("CSV", "*.csv"),
                       ("Alla", "*.*"),
                   ]
               ) or []]).grid(row=0, column=1, sticky="ew", padx=(0,8), pady=(8,4))
    ttk.Button(lf_in, text="Rensa listan", command=lambda: files_list.delete(0, tk.END))\
        .grid(row=1, column=1, sticky="ew", padx=(0,8))

    # Utdata + basnamn
    lf_out = ttk.LabelFrame(main, text="Utdatakatalog och basnamn"); lf_out.pack(fill="x", expand=False, pady=(0,10))
    var_outdir = tk.StringVar(value=""); var_basename = tk.StringVar(value=default_basename)
    ttk.Label(lf_out, text="Basnamn (utan ändelse):").grid(row=0, column=0, sticky="w", padx=8, pady=(8,4))
    ttk.Entry(lf_out, textvariable=var_basename).grid(row=0, column=1, sticky="ew", padx=8, pady=(8,4))
    ttk.Label(lf_out, text="Mapp för 'Results/':").grid(row=1, column=0, sticky="w", padx=8)
    ttk.Entry(lf_out, textvariable=var_outdir).grid(row=1, column=1, sticky="ew", padx=8)
    ttk.Button(lf_out, text="Välj mapp…",
               command=lambda: var_outdir.set(filedialog.askdirectory(title="Välj mapp") or var_outdir.get()))\
        .grid(row=1, column=2, sticky="ew", padx=(0,8))
    lf_out.columnconfigure(1, weight=1)

    # Diagram
    lf_plot = ttk.LabelFrame(main, text="Diagram (valfritt)"); lf_plot.pack(fill="x", expand=False, pady=(0,10))
    var_do_plots_summary = tk.BooleanVar(value=True)
    var_do_plots_pernight = tk.BooleanVar(value=True)
    var_generate_html = tk.BooleanVar(value=False)
    ttk.Checkbutton(lf_plot, text="Generera samlade diagram", variable=var_do_plots_summary)\
        .grid(row=0, column=0, sticky="w", padx=8, pady=(8,4))
    ttk.Checkbutton(lf_plot, text="Generera natt-för-natt diagram", variable=var_do_plots_pernight)\
        .grid(row=0, column=1, sticky="w", padx=8, pady=(8,4))
    ttk.Checkbutton(lf_plot, text="Generera interaktiv HTML (Parallel Bat Graph)", variable=var_generate_html)\
        .grid(row=0, column=2, sticky="w", padx=8, pady=(8,4))

    var_time_mode = tk.StringVar(value="manual")
    var_tstart = tk.StringVar(value="21:00"); var_tend = tk.StringVar(value="04:30")
    ttk.Radiobutton(lf_plot, text="X-axel: automatisk (från data)", value="auto", variable=var_time_mode)\
        .grid(row=1, column=0, sticky="w", padx=8)
    ttk.Radiobutton(lf_plot, text="X-axel: eget intervall", value="manual", variable=var_time_mode)\
        .grid(row=1, column=1, sticky="w")
    ttk.Label(lf_plot, text="Starttid (HH:MM):").grid(row=2, column=0, sticky="e", padx=8, pady=(4,8))
    ent_start = ttk.Entry(lf_plot, textvariable=var_tstart, width=10); ent_start.grid(row=2, column=1, sticky="w", pady=(4,8))
    ttk.Label(lf_plot, text="Sluttid (HH:MM):").grid(row=2, column=2, sticky="e", padx=8, pady=(4,8))
    ent_end = ttk.Entry(lf_plot, textvariable=var_tend, width=10); ent_end.grid(row=2, column=3, sticky="w", pady=(4,8))

    var_ymax_mode = tk.StringVar(value="fixed")
    ttk.Label(lf_plot, text="Y-axel läge:").grid(row=3, column=0, sticky="e", padx=8, pady=(0,8))
    ttk.Radiobutton(lf_plot, text="Fast (gemensam Y-max)", value="fixed", variable=var_ymax_mode)\
        .grid(row=3, column=1, sticky="w")
    ttk.Radiobutton(lf_plot, text="Zoomad (data-max + 10%)", value="zoomed", variable=var_ymax_mode)\
        .grid(row=3, column=2, sticky="w")

    ttk.Label(lf_plot, text="Bas-mapp för diagram (valfritt):").grid(row=4, column=0, sticky="e", padx=8, pady=(0,8))
    var_diagdir = tk.StringVar(value="")
    ttk.Entry(lf_plot, textvariable=var_diagdir).grid(row=4, column=1, sticky="ew", padx=8, pady=(0,8), columnspan=2)
    ttk.Button(lf_plot, text="Välj…", command=lambda: var_diagdir.set(filedialog.askdirectory(title="Välj basmapp") or var_diagdir.get()))\
        .grid(row=4, column=3, sticky="ew", pady=(0,8))
    lf_plot.columnconfigure(1, weight=1)

    # Färger + paleta (#8 SOF support)
    lf_colors = ttk.LabelFrame(main, text="Färger (valfritt – lämna tomt för standard)"); lf_colors.pack(fill="x", expand=False)
    def bind_preview(var, preview_widget):
        def _cb(*_):
            hx = _validate_hex(var.get())
            preview_widget.config(background=(hx or "#FFFFFF"))
        var.trace_add("write", _cb); _cb()
    def color_row(row, label, vars_tuple):
        ttk.Label(lf_colors, text=label).grid(row=row, column=0, sticky="w", padx=8, pady=(8,2))
        def one(col_ix, var):
            e = ttk.Entry(lf_colors, textvariable=var, width=9)
            e.grid(row=row, column=1 + col_ix*3, padx=(2,2), pady=(8,2), sticky="w")
            prev = tk.Label(lf_colors, text="  ", width=2, relief="groove")
            prev.grid(row=row, column=2 + col_ix*3, padx=(0,2), pady=(8,2), sticky="w")
            bind_preview(var, prev)
            def choose():
                init = _validate_hex(var.get()) or "#FFFFFF"
                _, hx = colorchooser.askcolor(color=init, title="Välj färg")
                if hx:
                    var.set(hx.upper())
            ttk.Button(lf_colors, text="Välj…", command=choose)\
                .grid(row=row, column=3 + col_ix*3, padx=(0,4), pady=(8,2), sticky="w")
        for i, v in enumerate(vars_tuple): one(i, v)
    var_nvi_soc=tk.StringVar(value=""); var_nvi_sof=tk.StringVar(value=""); var_nvi_fodo=tk.StringVar(value=""); var_nvi_forbi=tk.StringVar(value="")
    var_art_soc=tk.StringVar(value=""); var_art_sof=tk.StringVar(value=""); var_art_fodo=tk.StringVar(value=""); var_art_forbi=tk.StringVar(value="")
    color_row(0, "NVI – SOC / SOF / FOD / FORBI:", (var_nvi_soc, var_nvi_sof, var_nvi_fodo, var_nvi_forbi))
    color_row(1, "ART – SOC / SOF / FOD / FORBI:", (var_art_soc, var_art_sof, var_art_fodo, var_art_forbi))
    for c in range(13): lf_colors.columnconfigure(c, weight=0)
    lf_colors.columnconfigure(12, weight=1)

    # Knappfält
    btns = ttk.Frame(main); btns.pack(fill="x", pady=(12,0))
    settings = {}; done = {"ok": False}
    def on_ok():
        files = list(files_list.get(0, tk.END))
        if not files:
            messagebox.showwarning("GUI", "Välj minst en indatafil."); return
        outdir = var_outdir.get().strip()
        if not outdir:
            messagebox.showwarning("GUI", "Välj mapp där 'Results/' ska skapas."); return
        basename = var_basename.get().strip() or DEFAULT_OUT_BASENAME

        tmode = var_time_mode.get()
        tstart, tend = var_tstart.get().strip(), var_tend.get().strip()
        custom_range = None
        if tmode == "manual":
            try:
                validate_hhmm(tstart); validate_hhmm(tend)
            except Exception as ex:
                messagebox.showerror("GUI", f"Felaktigt tidsformat. Använd HH:MM. ({ex})"); return
            custom_range = (tstart, tend)

        settings.update({
            "input_files": files,
            "base_dir": outdir,
            "base_name": basename,
            "custom_time_range": custom_range,
            "ymax_mode": var_ymax_mode.get(),
            "diagram_base": (var_diagdir.get().strip() or None),
            "do_plots_summary": bool(var_do_plots_summary.get()),
            "do_plots_pernight": bool(var_do_plots_pernight.get()),
            "generate_html": bool(var_generate_html.get()),
            "colors": {
                "NVI": {
                    "Socialt": _validate_hex(var_nvi_soc.get()),
                    "Socialt - läte": _validate_hex(var_nvi_soc.get()),
                    "Socialt - flyg": _validate_hex(var_nvi_sof.get()),
                    "Födosökande": _validate_hex(var_nvi_fodo.get()),
                    "Förbiflygande": _validate_hex(var_nvi_forbi.get()),
                },
                "ART": {
                    "Socialt": _validate_hex(var_art_soc.get()),
                    "Socialt - läte": _validate_hex(var_art_soc.get()),
                    "Socialt - flyg": _validate_hex(var_art_sof.get()),
                    "Födosökande": _validate_hex(var_art_fodo.get()),
                    "Förbiflygande": _validate_hex(var_art_forbi.get()),
                },
            }
        })
        done["ok"] = True; root.destroy()
    def on_cancel():
        root.destroy(); sys.exit(0)
    ttk.Button(btns, text="Avbryt", command=on_cancel).pack(side="right")
    ttk.Button(btns, text="Starta", command=on_ok).pack(side="right", padx=(0,8))
    root.mainloop()
    if not done["ok"]: sys.exit(0)
    return settings

# ================== Start: hämta inställningar ==================
def run_analysis(settings: dict):
    input_files = settings["input_files"]
    base_dir    = settings["base_dir"]
    base_name   = settings["base_name"]

    report_progress(20, "Reading source tables")

    if os.path.basename(base_dir).lower() == "results":
        results_dir = base_dir
    else:
        results_dir = os.path.join(base_dir, "results")

    combined_dir = os.path.join(results_dir, "combined")
    inputs_dir = settings["diagram_base"] if settings.get("diagram_base") else results_dir
    os.makedirs(combined_dir, exist_ok=True)
    os.makedirs(inputs_dir, exist_ok=True)

    out_path_nvi = os.path.join(combined_dir, f"{base_name}_NVI.xlsx")
    out_path_art = os.path.join(combined_dir, f"{base_name}_ART.xlsx")

    # Resolve dataset SOF presence and class palette config
    class_cfg = resolve_dataset_class_config(input_files, settings.get("colors"))
    type_order_overview = class_cfg["type_order"]
    code_map = class_cfg["code_map"]

    def scheme_table_nvi():
        fills = {
            typ: fill_from_hex(col) for typ, col in class_cfg["nvi_color_dict"].items()
        }
        fills["Fodosökande"] = fill_from_hex(class_cfg["nvi_color_dict"].get("Födosökande", HEX_NVI_FODO))
        fills["Förbiflygande"] = fill_from_hex(HEX_TABLE_FORBI)
        return fills

    def scheme_table_art():
        fills = {
            typ: fill_from_hex(col) for typ, col in class_cfg["art_color_dict"].items()
        }
        fills["Fodosökande"] = fill_from_hex(class_cfg["art_color_dict"].get("Födosökande", HEX_ART_FODO))
        fills["Förbiflygande"] = fill_from_hex(HEX_TABLE_FORBI)
        return fills

    unique_stems = get_unique_input_stems(input_files)

    # ================== STEG 1: Bygg Excel-översikt ==================
    used_sheet_names = set()
    sheets_to_write = []
    counts_per_file = {}
    total_ljud_per_file = {}
    nights_per_file = {}
    all_species_latin = set()

    for path in input_files:
        sheet_name = safe_sheet_name(path, used_sheet_names)
        df = read_input_table(path)
        sheets_to_write.append((sheet_name, df))
        total_ljud_per_file[sheet_name] = int(len(df))  # inkl. Noise
        nights_per_file[sheet_name] = count_nights(df)

        if "MANUAL ID" not in df.columns:
            counts_per_file[sheet_name] = {}
            continue

        df["__list"] = df["MANUAL ID"].map(extract_species_and_type)
        long = df.explode("__list"); long = long[long["__list"].notna()]
        if long.empty:
            counts_per_file[sheet_name] = {}
            continue

        long[["ArtLatin", "Beteendetyper_code"]] = pd.DataFrame(long["__list"].tolist(), index=long.index)
        long["Beteendetyper"] = long["Beteendetyper_code"].map(code_map)
        long = long[long["ArtLatin"].astype(str).str.strip().str.lower() != "noise"]
        all_species_latin.update(long["ArtLatin"].astype(str).str.strip().tolist())
        grp = long.groupby(["ArtLatin", "Beteendetyper"]).size()
        counts_per_file[sheet_name] = {(sp, typ): int(n) for (sp, typ), n in grp.items()}

    species_sorted_latin = sorted(all_species_latin, key=species_sort_key)
    file_cols = list(counts_per_file.keys())

    rows_data = []
    for latin in species_sorted_latin:
        disp = display_label_multiline(latin)
        for typ in type_order_overview:
            row = {"Art": disp, "Beteendetyper": typ}
            for col in file_cols:
                val = counts_per_file.get(col, {}).get((latin, typ), 0)
                row[col] = ("" if val == 0 else int(val))
            rows_data.append(row)

    sum_row = {"Art": "", "Beteendetyper": "Fladdermusregistreringar"}
    for col in file_cols: sum_row[col] = int(sum(counts_per_file.get(col, {}).values()))
    rows_data.append(sum_row)

    nights_row = {"Art": "", "Beteendetyper": "Antal nätter"}
    for col in file_cols:
        n = nights_per_file.get(col)
        nights_row[col] = ("" if not n else int(n))
    rows_data.append(nights_row)

    per_night_row = {"Art": "", "Beteendetyper": "Antal registreringar / natt"}
    for col in file_cols: per_night_row[col] = ""
    rows_data.append(per_night_row)

    tot_row = {"Art": "", "Beteendetyper": "Total antal ljud"}
    for col in file_cols: tot_row[col] = int(total_ljud_per_file.get(col, 0))
    rows_data.append(tot_row)

    overview_df = pd.DataFrame(rows_data, columns=["Art", "Beteendetyper"] + file_cols)
    num_species = len(species_sorted_latin)

    def write_overview_to(path_out):
        with pd.ExcelWriter(path_out, engine="openpyxl") as writer:
            overview_df.to_excel(writer, sheet_name="Översikt", index=False)
            for sheet_name, df_orig in sheets_to_write:
                df_orig.to_excel(writer, sheet_name=sheet_name, index=False)

    def format_overview(path_out, scheme_fills, num_species_rows, file_cols_list):
        wb = load_workbook(path_out); ws = wb["Översikt"]
        max_row = ws.max_row; max_col = ws.max_column
        header_row = 1; data_start = header_row + 1
        num_types_per_sp = len(type_order_overview)
        num_species_rows_total = num_species_rows * num_types_per_sp
        sum_row_idx      = data_start + num_species_rows_total
        nights_row_idx   = sum_row_idx + 1
        pernight_row_idx = nights_row_idx + 1
        total_row_idx    = pernight_row_idx + 1

        # rubriker
        for c in range(1, max_col + 1):
            cell = ws.cell(row=header_row, column=c)
            cell.fill = FILL_HDR; cell.font = Font(color=hex_to_argb(HEX_HEADER_FG), bold=True)
            cell.alignment = Alignment(vertical="center", horizontal="center")
        ws.row_dimensions[header_row].height = 18

        # kolumnbredder
        ws.column_dimensions["A"].width = 44; ws.column_dimensions["B"].width = 24
        for idx, col_name in enumerate(file_cols_list, start=3):
            header_text = str(col_name)
            ws.column_dimensions[get_column_letter(idx)].width = max(12, min(50, int(len(header_text) * 1.1)))

        # slå ihop artetiketter
        if num_species_rows_total > 0:
            current = data_start; merge_end_limit = sum_row_idx - 1
            while current <= merge_end_limit:
                art_val = ws[f"A{current}"].value
                if not art_val: current += 1; continue
                end = current
                while end + 1 <= merge_end_limit and ws[f"A{end+1}"].value == art_val:
                    end += 1
                if end > current:
                    ws.merge_cells(start_row=current, start_column=1, end_row=end, end_column=1)
                ws.cell(row=current, column=1).alignment = Alignment(vertical="center", wrap_text=True)
                current = end + 1

        # färg rader per typ
        for r in range(data_start, sum_row_idx):
            typ = ws.cell(row=r, column=2).value
            fill = scheme_fills.get(typ)
            if fill:
                ws.cell(row=r, column=2).fill = fill
                for c in range(3, max_col + 1):
                    val = ws.cell(row=r, column=c).value
                    if val not in (None, "", 0): ws.cell(row=r, column=c).fill = fill

        # summeringsrader (bold)
        for r in (sum_row_idx, nights_row_idx, pernight_row_idx, total_row_idx):
            for c in range(1, max_col + 1):
                ws.cell(row=r, column=c).font = Font(bold=True)
                ws.cell(row=r, column=c).alignment = Alignment(vertical="center")

        # formel
        for col_idx in range(3, max_col + 1):
            L = get_column_letter(col_idx)
            cell = ws.cell(row=pernight_row_idx, column=col_idx)
            cell.value = f'=IFERROR({L}{sum_row_idx}/{L}{nights_row_idx},"")'
            cell.number_format = "0.0"

        # blocklinjer
        for i in range(num_species_rows):
            top_row = data_start + i * num_types_per_sp
            for c in range(1, max_col + 1):
                old = ws.cell(row=top_row, column=c).border
                ws.cell(row=top_row, column=c).border = Border(left=old.left, right=old.right, top=BORDER_MEDIUM, bottom=old.bottom)
        for c in range(1, max_col + 1):
            old = ws.cell(row=sum_row_idx, column=c).border
            ws.cell(row=sum_row_idx, column=c).border = Border(left=old.left, right=old.right, top=BORDER_MEDIUM, bottom=old.bottom)
            old = ws.cell(row=total_row_idx, column=c).border
            ws.cell(row=total_row_idx, column=c).border = Border(left=old.left, right=BORDER_MEDIUM, top=old.top, bottom=BORDER_MEDIUM)

        # vertikala avgränsningar
        for c in range(3, max_col + 1):
            for r in range(header_row, max_row + 1):
                old = ws.cell(row=r, column=c).border
                ws.cell(row=r, column=c).border = Border(left=BORDER_MEDIUM, right=old.right, top=old.top, bottom=old.bottom)
        for r in range(header_row, max_row + 1):
            oldA = ws.cell(row=r, column=1).border
            ws.cell(row=r, column=1).border = Border(left=BORDER_MEDIUM, right=oldA.right, top=oldA.top, bottom=oldA.bottom)
            oldB = ws.cell(row=r, column=2).border
            ws.cell(row=r, column=2).border = Border(left=BORDER_MEDIUM, right=oldB.right, top=oldB.top, bottom=oldB.bottom)
            oldR = ws.cell(row=r, column=max_col).border
            ws.cell(row=r, column=max_col).border = Border(left=oldR.left, right=BORDER_MEDIUM, top=oldR.top, bottom=BORDER_MEDIUM)

        wb.save(path_out); wb.close()

    # Skriv båda filer
    write_overview_to(out_path_nvi); format_overview(out_path_nvi, scheme_table_nvi(), num_species, file_cols)
    print(f"Klar! Sparad fil (NVI): {out_path_nvi}")
    write_overview_to(out_path_art); format_overview(out_path_art, scheme_table_art(), num_species, file_cols)
    print(f"Klar! Sparad fil (ART): {out_path_art}")

    report_progress(40, "Generating combined NVI/ART workbooks")

    # Interactive HTML Parallel Bat Graph (optional)
    if settings.get("generate_html"):
        report_progress(55, "Generating interactive HTML visualization")
        if os.environ.get("APP_ENV") == "testing" and os.environ.get("TEST_FAIL_HTML_GEN") == "1":
            raise RuntimeError("Simulated HTML generation failure")
        try:
            from parallel_graph.export import generate_outputs
            from parallel_graph.models import SourceSpec
        except ModuleNotFoundError:
            pg_src = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "parallel-graph", "src"))
            if pg_src not in sys.path:
                sys.path.insert(0, pg_src)
            from parallel_graph.export import generate_outputs
            from parallel_graph.models import SourceSpec

        from pathlib import Path
        sources = [SourceSpec(path=Path(f), name=unique_stems[f]) for f in input_files]
        custom_time_range = settings.get("custom_time_range")
        generate_outputs(
            sources=sources,
            output_directory=Path(combined_dir),
            language="pl",
            time_range=custom_time_range,
        )
        print(f"Klar! Sparade parallel HTML, CSV och report i {combined_dir}")

    if settings.get("open_files", True):
        open_file(out_path_nvi); open_file(out_path_art)

    # ================== Uruchom tworzenie wykresów ==================
    if settings.get("do_plots_summary") or settings.get("do_plots_pernight"):
        print(f"Resultat kommer att sparas i: {inputs_dir}")

        if settings.get("do_plots_summary"):
            report_progress(70, "Generating summary static charts")
            generate_summary_diagrams(
                input_files_list=input_files,
                inputs_dir=inputs_dir,
                unique_stems=unique_stems,
                custom_time_range=settings.get("custom_time_range"),
                ymax_mode=settings.get("ymax_mode", "fixed"),
                class_cfg=class_cfg,
            )
        if settings.get("do_plots_pernight"):
            report_progress(85, "Generating per-night static charts")
            generate_pernight_diagrams(
                input_files_list=input_files,
                inputs_dir=inputs_dir,
                unique_stems=unique_stems,
                custom_time_range=settings.get("custom_time_range"),
                ymax_mode=settings.get("ymax_mode", "fixed"),
                class_cfg=class_cfg,
            )

    report_progress(95, "Finalizing charts and output tables")


# ================== STEG 2: Diagram (linje + stapel) ==================
def _plot_for_subset(df_subset, custom_time_range, y_lim_global,
                     out_lines, out_stacks_art, out_stacks_nvi,
                     ymax_mode: str, class_cfg: dict,
                     night_label: str | None = None):
    df = df_subset.copy()
    if "MANUAL ID" not in df.columns: return

    time_col = detect_column(df, ["time", "tid"])
    if time_col is None: return

    if custom_time_range:
        start_str, stop_str = custom_time_range
        df = df[df[time_col].apply(lambda t: is_time_in_range(t, start_str, stop_str))]
        if df.empty: return

    df["species_type_list"] = df["MANUAL ID"].map(extract_species_and_type)

    def time_to_interval(val):
        hm = _hm_from_any(val)
        if hm is None: return ""
        h, m = hm; minutes = int((m // 15) * 15)
        return f"{h:02d}:{minutes:02d}"
    df["interval"] = df[time_col].map(time_to_interval)

    df_long = df.explode("species_type_list")
    df_long = df_long[df_long["species_type_list"].notna()]
    if df_long.empty: return
    df_long[["species", "obs_code"]] = pd.DataFrame(df_long["species_type_list"].tolist(), index=df_long.index)

    code_map = class_cfg["code_map"]
    type_order = class_cfg["type_order"]
    colors_art = class_cfg["colors_art"]
    colors_nvi = class_cfg["colors_nvi"]
    art_color_dict = class_cfg["art_color_dict"]
    nvi_color_dict = class_cfg["nvi_color_dict"]

    df_long["obs_type"] = df_long["obs_code"].map(code_map)
    df_long = df_long[df_long["species"].astype(str).str.strip().str.lower() != "noise"]
    df_long = df_long[df_long["species"].astype(str).str.strip() != ""]
    if df_long.empty: return

    # tidsintervall
    if custom_time_range:
        all_intervals = build_manual_interval_sequence(custom_time_range)
    else:
        dt_series = df[time_col].apply(str_to_dt).dropna()
        if len(dt_series) == 0:
            ints = [s for s in df["interval"].astype(str).tolist() if s and s.lower() != "nan"]
            dt_from_int = [interval_to_sortkey(s) for s in ints]
            dt_from_int = [d for d in dt_from_int if d is not None]
            if not dt_from_int: return
            min_dt = round_down_15(min(dt_from_int)); max_dt = round_up_15(max(dt_from_int))
        else:
            min_dt = round_down_15(min(dt_series)); max_dt = round_up_15(max(dt_series))

        all_intervals = []
        t = min_dt
        while t <= max_dt:
            all_intervals.append(t.strftime("%H:%M")); t += timedelta(minutes=15)
        all_intervals = list(dict.fromkeys(all_intervals))

    # Agg
    agg = df_long.groupby(["interval", "species", "obs_type"]).size().reset_index(name="antal")
    agg["interval"] = pd.Categorical(agg["interval"], categories=all_intervals, ordered=True)
    species_list = sorted(df_long["species"].unique(), key=species_sort_key)

    os.makedirs(out_lines, exist_ok=True)
    os.makedirs(out_stacks_art, exist_ok=True)
    os.makedirs(out_stacks_nvi, exist_ok=True)

    # LINJE – samlingsdiagram
    agg_line = df_long.groupby(["interval", "species"]).size().reset_index(name="antal")
    agg_line["interval"] = pd.Categorical(agg_line["interval"], categories=all_intervals, ordered=True)
    pivot_line = agg_line.pivot(index="interval", columns="species", values="antal").fillna(0)
    pivot_line = pivot_line.reindex(all_intervals, fill_value=0)
    total_obs = df_long.shape[0]

    ymax_line_all = compute_ymax_with_headroom(pivot_line.max().max(), headroom_factor=1.10) if ymax_mode == "zoomed" else y_lim_global

    plt.figure(figsize=(12, 6))
    ax_all = plt.gca()
    pivot_line.plot(ax=ax_all, marker='o')
    ax_all.set_xlabel("Tid (15-minutersintervall)"); ax_all.set_ylabel("Antal ljudfiler")
    title_all = "Fladdermusobservationer – alla arter"
    if night_label: title_all += f" – natt {night_label}"
    title_all += f", antal observerade beteenden: {total_obs}"
    ax_all.set_title(title_all)
    ax_all.legend(title="Art", bbox_to_anchor=(1.05, 1), loc='upper left')
    ax_all.yaxis.set_major_locator(MaxNLocator(integer=True))
    ax_all.set_xticks(range(len(pivot_line.index))); ax_all.set_xticklabels(pivot_line.index, rotation=270)
    ax_all.set_ylim(0, ymax_line_all)
    plt.tight_layout(); plt.grid(True, axis='y')
    plt.savefig(os.path.join(out_lines, "alla_arter.png")); plt.close()

    # LINJE – per art
    for species in species_list:
        total_sp = int(pivot_line[species].sum())
        ymax_line_sp = compute_ymax_with_headroom(pivot_line[species].max(), headroom_factor=1.10) if ymax_mode == "zoomed" else y_lim_global
        plt.figure(figsize=(10, 4)); ax = plt.gca()
        ax.plot(pivot_line.index, pivot_line[species], marker='o')
        ax.set_xlabel("Tid (15-minutersintervall)"); ax.set_ylabel("Antal ljudfiler")
        ax.set_title(format_title(species, total_sp, night_label))
        ax.yaxis.set_major_locator(MaxNLocator(integer=True))
        ax.set_xticks(range(len(pivot_line.index))); ax.set_xticklabels(pivot_line.index, rotation=270)
        ax.set_ylim(0, ymax_line_sp)
        plt.grid(True, axis='y'); plt.tight_layout()
        plt.savefig(os.path.join(out_lines, f"{safe_filename(species)}.png")); plt.close()

    # STAPEL – ART per art
    for species in species_list:
        plot_data = (
            agg[agg["species"] == species]
            .pivot(index="interval", columns="obs_type", values="antal")
            .fillna(0)
            .reindex(all_intervals, fill_value=0)
            .reindex(columns=type_order, fill_value=0)
        )
        ymax_bar_sp = compute_ymax_with_headroom(plot_data.sum(axis=1).max(), headroom_factor=1.10) if ymax_mode == "zoomed" else y_lim_global
        ax = plot_data.plot(kind="bar", stacked=True, color=colors_art, figsize=(14, 6))
        plt.xlabel("Tid (15-minutersintervall)"); plt.ylabel("Antal ljudfiler")
        plt.title(format_title(species, int(plot_data.values.sum()), night_label))
        plt.legend(title="Beteendetyper")
        ax.yaxis.set_major_locator(MaxNLocator(integer=True)); ax.set_ylim(0, ymax_bar_sp)
        plt.xticks(rotation=270); plt.tight_layout(); plt.grid(True, axis='y')
        plt.savefig(os.path.join(out_stacks_art, f"{safe_filename(species)}.png")); plt.close()

    # STAPEL – NVI per art
    for species in species_list:
        plot_data = (
            agg[agg["species"] == species]
            .pivot(index="interval", columns="obs_type", values="antal")
            .fillna(0)
            .reindex(all_intervals, fill_value=0)
            .reindex(columns=type_order, fill_value=0)
        )
        ymax_bar_sp = compute_ymax_with_headroom(plot_data.sum(axis=1).max(), headroom_factor=1.10) if ymax_mode == "zoomed" else y_lim_global
        ax = plot_data.plot(kind="bar", stacked=True, color=colors_nvi, figsize=(14, 6))
        plt.xlabel("Tid (15-minutersintervall)"); plt.ylabel("Antal ljudfiler")
        plt.title(format_title(species, int(plot_data.values.sum()), night_label))
        plt.legend(title="Beteendetyper")
        ax.yaxis.set_major_locator(MaxNLocator(integer=True)); ax.set_ylim(0, ymax_bar_sp)
        plt.xticks(rotation=270); plt.tight_layout(); plt.grid(True, axis='y')
        plt.savefig(os.path.join(out_stacks_nvi, f"{safe_filename(species)}.png")); plt.close()

    # ALL-SPECIES GROUPED STACKED CHART (#9)
    title_all_stacked = "Fladdermusobservationer – alla arter"
    if night_label: title_all_stacked += f" – natt {night_label}"
    title_all_stacked += f", antal observerade beteenden: {total_obs}"

    # ART version
    core._plot_all_species_grouped_stacked(
        df_long=df_long,
        all_intervals=all_intervals,
        species_list=species_list,
        type_order=type_order,
        color_dict=art_color_dict,
        out_path=os.path.join(out_stacks_art, "alla_arter.png"),
        title_text=title_all_stacked,
        ymax_mode=ymax_mode,
        y_lim_global=y_lim_global,
    )

    # NVI version
    core._plot_all_species_grouped_stacked(
        df_long=df_long,
        all_intervals=all_intervals,
        species_list=species_list,
        type_order=type_order,
        color_dict=nvi_color_dict,
        out_path=os.path.join(out_stacks_nvi, "alla_arter.png"),
        title_text=title_all_stacked,
        ymax_mode=ymax_mode,
        y_lim_global=y_lim_global,
    )

def generate_summary_diagrams(input_files_list, inputs_dir, unique_stems, custom_time_range, ymax_mode="fixed", class_cfg=None):
    if class_cfg is None:
        class_cfg = resolve_dataset_class_config(input_files_list)
    y_lim = compute_global_ymax_across_files(input_files_list, custom_time_range, code_map=class_cfg["code_map"])
    print(f"Global gemensam Y-max (SAMLADE): {y_lim} (mode: {ymax_mode})")
    for input_file in input_files_list:
        stem = unique_stems[input_file]
        df_full = read_input_table(input_file)
        stem_dir = os.path.join(inputs_dir, stem)
        out_lines = os.path.join(stem_dir, "summary", "line")
        out_stacks_art = os.path.join(stem_dir, "summary", "stacked_ART")
        out_stacks_nvi = os.path.join(stem_dir, "summary", "stacked_NVI")
        os.makedirs(out_lines, exist_ok=True)
        os.makedirs(out_stacks_art, exist_ok=True)
        os.makedirs(out_stacks_nvi, exist_ok=True)
        _plot_for_subset(df_full, custom_time_range, y_lim, out_lines, out_stacks_art, out_stacks_nvi,
                         ymax_mode=ymax_mode, class_cfg=class_cfg, night_label=None)

def generate_pernight_diagrams(input_files_list, inputs_dir, unique_stems, custom_time_range, ymax_mode="fixed", class_cfg=None):
    if class_cfg is None:
        class_cfg = resolve_dataset_class_config(input_files_list)
    y_lim = compute_global_ymax_across_files_and_nights(input_files_list, custom_time_range, code_map=class_cfg["code_map"])
    print(f"Global gemensam Y-max (NATT-FÖR-NATT): {y_lim} (mode: {ymax_mode})")
    for input_file in input_files_list:
        stem = unique_stems[input_file]
        df_full = read_input_table(input_file)
        date_col = detect_column(df_full, ["date", "datum"])
        time_col = detect_column(df_full, ["time", "tid"])
        if not date_col:
            stem_dir = os.path.join(inputs_dir, stem)
            out_lines = os.path.join(stem_dir, "nights", "utan_datum", "line")
            out_stacks_art = os.path.join(stem_dir, "nights", "utan_datum", "stacked_ART")
            out_stacks_nvi = os.path.join(stem_dir, "nights", "utan_datum", "stacked_NVI")
            os.makedirs(out_lines, exist_ok=True)
            os.makedirs(out_stacks_art, exist_ok=True)
            os.makedirs(out_stacks_nvi, exist_ok=True)
            _plot_for_subset(df_full, custom_time_range, y_lim, out_lines, out_stacks_art, out_stacks_nvi,
                             ymax_mode=ymax_mode, class_cfg=class_cfg, night_label=None)
            continue
        df_full["__night"] = df_full.apply(lambda r: row_night_key(r[date_col], r[time_col] if time_col in r else None), axis=1)
        for night_start, sub in df_full.groupby("__night"):
            if night_start is None or sub.empty: continue
            night_lab = night_label_str(night_start)
            night_folder = night_start.isoformat()
            stem_dir = os.path.join(inputs_dir, stem)
            out_lines = os.path.join(stem_dir, "nights", night_folder, "line")
            out_stacks_art = os.path.join(stem_dir, "nights", night_folder, "stacked_ART")
            out_stacks_nvi = os.path.join(stem_dir, "nights", night_folder, "stacked_NVI")
            os.makedirs(out_lines, exist_ok=True)
            os.makedirs(out_stacks_art, exist_ok=True)
            os.makedirs(out_stacks_nvi, exist_ok=True)
            _plot_for_subset(sub, custom_time_range, y_lim, out_lines, out_stacks_art, out_stacks_nvi,
                             ymax_mode=ymax_mode, class_cfg=class_cfg, night_label=night_lab)

if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="Merge bat tables and create charts.")
    parser.add_argument("--headless", action="store_true", help="Run without GUI")
    parser.add_argument("--input-files", nargs="+", help="Input XLSX or CSV files")
    parser.add_argument("--base-dir", help="Base output directory")
    parser.add_argument("--base-name", default=DEFAULT_OUT_BASENAME, help="Output base name")
    parser.add_argument("--no-plots-summary", action="store_true", help="Skip summary plots")
    parser.add_argument("--no-plots-pernight", action="store_true", help="Skip per-night plots")
    parser.add_argument("--time-mode", choices=["manual", "auto"], default="manual", help="Time range mode")
    parser.add_argument("--time-start", default="21:00", help="Custom start time HH:MM")
    parser.add_argument("--time-end", default="04:30", help="Custom end time HH:MM")
    parser.add_argument("--ymax-mode", "--y-axis-mode", choices=["fixed", "zoomed"], default="fixed", help="Y-axis mode (fixed or zoomed)")
    parser.add_argument("--enable-html", action="store_true", help="Generate interactive HTML parallel bat graph")
    parser.add_argument("--no-open", action="store_true", help="Do not open files after processing")

    args, _ = parser.parse_known_args()

    if args.headless or args.input_files:
        custom_time_range = None
        if args.time_mode == "manual":
            tstart = args.time_start or "21:00"
            tend = args.time_end or "04:30"
            validate_hhmm(tstart)
            validate_hhmm(tend)
            custom_time_range = (tstart, tend)

        settings = {
            "input_files": [os.path.abspath(f) for f in (args.input_files or [])],
            "base_dir": os.path.abspath(args.base_dir) if args.base_dir else os.getcwd(),
            "base_name": args.base_name,
            "custom_time_range": custom_time_range,
            "ymax_mode": args.ymax_mode,
            "do_plots_summary": not args.no_plots_summary,
            "do_plots_pernight": not args.no_plots_pernight,
            "generate_html": args.enable_html,
            "colors": {
                "NVI": {"Socialt": None, "Födosökande": None, "Förbiflygande": None},
                "ART": {"Socialt": None, "Födosökande": None, "Förbiflygande": None},
            },
            "open_files": not args.no_open and not args.headless,
        }
        run_analysis(settings)
    else:
        settings = gui_collect_settings()
        run_analysis(settings)

